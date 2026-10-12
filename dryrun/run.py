"""End-to-end dry run of pipeline-factory against local stand-ins (see dryrun/servers.py and mock_fivetran.py).

Drives the `pf` CLI exactly as a user and CI would, runs every built connector in the official Fivetran local
tester (`fivetran debug`, DuckDB output), and asserts what lands. Writes <state>/DRY-RUN-REPORT.json.

    python dryrun/run.py <state_dir> [--llm replay|record] [--only sftp,ftp,rds,deploy,drills]

The LLM steps run through the real Claude Agent SDK in `record` mode (needs credentials) or from the recorded
fixtures in tests/llm-fixtures in `replay` mode (offline, deterministic).
"""
from __future__ import annotations

import argparse
import base64
import importlib.util
import json
import os
import shutil
import subprocess
import sys
import time
import traceback
from pathlib import Path

FRAMEWORK = Path(__file__).resolve().parents[1]
BIN = Path(sys.executable).parent
PF = str(BIN / "pf")
FIVETRAN = str(BIN / "fivetran")
OLD = time.time() - 3600  # deliveries staged "an hour ago" so readiness (min_age) passes


class Report:
    def __init__(self) -> None:
        self.checks: list[dict] = []
        self.scenario = ""

    def check(self, ok: bool, what: str, detail=None) -> bool:
        self.checks.append({"scenario": self.scenario, "check": what, "ok": bool(ok), "detail": detail})
        print(("  PASS " if ok else "  FAIL ") + what + ("" if ok or detail is None else f"  -> {str(detail)[:400]}"), flush=True)
        return ok


R = Report()


# ------------------------------------------------------------------------------------------------ helpers
def run(cmd: list[str], cwd: Path, env: dict | None = None, timeout: int = 900) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, cwd=cwd, env=env, capture_output=True, text=True, timeout=timeout)


def base_env(state: Path, llm: str, human: bool = True) -> dict:
    env = dict(os.environ)
    env.pop("CLAUDECODE", None)
    env.pop("PF_AGENT_SESSION", None)
    if not human:
        env["PF_AGENT_SESSION"] = "1"
    env["PF_LLM_PROVIDER"] = llm
    env["PF_LLM_FIXTURES"] = str(FRAMEWORK / "tests" / "llm-fixtures")
    env["HOME"] = str(state / "home")
    return env


def pf(repo: Path, *args: str, env: dict, ok: bool | None = True) -> dict:
    p = run([PF, "--json", *args], repo, env)
    try:
        out = json.loads(p.stdout)
    except json.JSONDecodeError:
        out = {"ok": False, "issues": [{"msg": (p.stdout + p.stderr)[-1500:]}], "data": {}}
    if ok is not None and out.get("ok") != ok:
        raise AssertionError(f"pf {' '.join(args)} -> ok={out.get('ok')}: {json.dumps(out.get('issues'))[:1500]}")
    return out


def pf_set(repo: Path, name: str, file: str, path: str, value, env: dict) -> None:
    pf(repo, "set", name, file, path, json.dumps(value) if not isinstance(value, str) or value.startswith(("{", "[")) else json.dumps(value), env=env, ok=None)


def fill(repo: Path, name: str, values: dict[tuple[str, str], object], env: dict) -> None:
    for (f, path), v in values.items():
        out = pf(repo, "set", name, f, path, json.dumps(v), env=env, ok=None)
        if not out.get("ok"):
            errs = [i for i in out.get("issues", []) if i.get("severity") == "error" and i.get("code") != "required"]
            if errs:
                raise AssertionError(f"set {f} {path}: {errs[:3]}")


def stage(tree: Path, dest: Path) -> list[str]:
    files = []
    for p in sorted(tree.rglob("*")):
        if p.is_file() and p.name != "expect.yaml":
            t = dest / p.relative_to(tree)
            t.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(p, t)
            os.utime(t, (OLD, OLD))
            files.append(str(p.relative_to(tree)))
    return files


def tester(connector: Path, config: dict, state: Path, label: str, env: dict) -> subprocess.CompletedProcess:
    cfg = state / f"config-{connector.parent.parent.name}.json"
    cfg.write_text(json.dumps(config))
    os.chmod(cfg, 0o600)
    p = run([FIVETRAN, "debug", "--configuration", str(cfg)], connector, env, timeout=900)
    (state / "logs").mkdir(exist_ok=True)
    (state / "logs" / f"{label}.log").write_text(p.stdout + "\n--- stderr ---\n" + p.stderr)
    return p


def wh(connector: Path, sql: str):
    import duckdb

    con = duckdb.connect(str(connector / "files" / "warehouse.db"), read_only=True)
    try:
        return con.execute(sql).fetchall()
    finally:
        con.close()


def s3(env_json: dict):
    import boto3

    return boto3.client("s3", region_name="us-east-1", endpoint_url=env_json["s3_endpoint"],
                        aws_access_key_id="testing", aws_secret_access_key="testing")


def load_builder():
    spec = importlib.util.spec_from_file_location("fx", FRAMEWORK / "tests" / "fixture_builders" / "oncology_labs.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def new_repo(state: Path, env: dict) -> Path:
    repo = state / "repo"
    if not repo.exists():
        repo.mkdir(parents=True)
        p = run([PF, "init-repo", str(repo)], state, env)
        assert p.returncode == 0, p.stderr
        run(["git", "init", "-q", "-b", "main"], repo)
        run(["git", "config", "user.email", "dryrun@example.com"], repo)
        run(["git", "config", "user.name", "dry run"], repo)
    return repo


# ------------------------------------------------------------------------------------------------ scenario: SFTP
def scenario_sftp(state: Path, E: dict, env: dict) -> None:
    R.scenario = "sftp"
    print("\n== SFTP: oncology_labs (zip packages, 10 structural checks, English rule, originals archive)", flush=True)
    repo = new_repo(state, env)
    shutil.copytree(FRAMEWORK / "pipelines" / "oncology_labs", repo / "pipelines" / "oncology_labs",
                    ignore=shutil.ignore_patterns("build", "intake"), dirs_exist_ok=True)
    fill(repo, "oncology_labs", {("source.yaml", "source.host"): "127.0.0.1", ("source.yaml", "source.port"): E["ports"]["sftp"],
                                 ("source.yaml", "source.host_key"): E["host_key"], ("source.yaml", "source.root"): "/outbound/oncology"}, env)
    v = pf(repo, "validate", "oncology_labs", env=env)
    R.check(v["ok"], "repointed spec validates (English rule stays accepted: only connection values changed)")
    b = pf(repo, "build", "oncology_labs", env=env)
    R.check(b["ok"], "pf build", b["data"])
    R.check(pf(repo, "build", "oncology_labs", "--check", env=env)["ok"], "pf build --check is deterministic")
    t = pf(repo, "test", "oncology_labs", env=env, ok=None)
    R.check(t["ok"], "pf test: 23 break-and-verify cases, every enforced check fires", t["issues"][:3])

    s = s3(E)
    if "az-rd-originals-dev" in {x["Name"] for x in s.list_buckets().get("Buckets", [])}:
        raise RuntimeError("stale originals bucket: restart dryrun/servers.py for a clean run")
    s.create_bucket(Bucket="az-rd-originals-dev", ObjectLockEnabledForBucket=True)
    s.put_object_lock_configuration(Bucket="az-rd-originals-dev", ObjectLockConfiguration={
        "ObjectLockEnabled": "Enabled", "Rule": {"DefaultRetention": {"Mode": "COMPLIANCE", "Days": 1}}})
    root = Path(E["sftp_root"]) / "outbound" / "oncology"
    staged = stage(repo / "pipelines" / "oncology_labs" / "fixtures" / "good", root)
    conn = repo / "pipelines" / "oncology_labs" / "build" / "connector"
    cfg = {"pf_environment": "dev", "sftp_private_key": Path(E["client_key"]).read_text(), "subject_key_salt": "dry-run-salt-1234567890",
           "originals_access_key_id_dev": "testing", "originals_secret_access_key_dev": "testing", "pf_originals_endpoint_url": E["s3_endpoint"]}

    p = tester(conn, cfg, state, "sftp-sync1", env)
    R.check(p.returncode == 0, "sync 1: fivetran debug exits 0", p.stdout[-800:] + p.stderr[-800:])
    rows = wh(conn, "select count(*), count(distinct subject_key) from tester.lab_results")[0]
    R.check(rows[0] == 54, "sync 1: 54 distinct lab results upserted (48 + 6 new keys; 6 re-sent keys updated)", rows)
    raw_ids = wh(conn, "select count(*) from tester.lab_results where subject_key like 'S0%'")[0][0]
    R.check(raw_ids == 0, "subject ids are pseudonymised (no clear vendor id in the lake)", raw_ids)
    aud = wh(conn, "select package, status from tester.pf_package_audit order by package")
    R.check([a[1] for a in aud] == ["loaded", "loaded"], "sync 1: both deliveries audited as loaded", aud)
    lin = wh(conn, "select count(*) from tester.lab_results where pf_package is null or pf_package_hash is null or pf_run_id is null")[0][0]
    R.check(lin == 0, "every row carries lineage (package, sha256, run id)", lin)
    objs = s.list_objects_v2(Bucket="az-rd-originals-dev").get("Contents", [])
    R.check(len(objs) == 2, "originals: 2 deliveries archived byte-for-byte (write-once, object lock)", [o["Key"] for o in objs])
    if objs:
        head = s.head_object(Bucket="az-rd-originals-dev", Key=objs[0]["Key"])
        R.check(head.get("ObjectLockMode") == "COMPLIANCE", "originals: archived object is under COMPLIANCE retention", head.get("ObjectLockMode"))

    fx = load_builder()
    d3 = fx.delivery(fx.rows(5, 1)[24:])  # new subject S0005, distinct bytes
    fx.put(root, "20240303/LB_ONC101_003.zip", d3, sidecar="0" * 64)
    d4 = fx.delivery([r.replace(",1,01MAR2024,03MAR2024,", ",4,04MAR2024,05MAR2024,") for r in fx.rows(1, 1)])
    fx.put(root, "20240304/LB_ONC101_004.zip", d4)
    for p_ in root.rglob("*"):
        os.utime(p_, (OLD, OLD))
    p = tester(conn, cfg, state, "sftp-sync2", env)
    R.check(p.returncode == 0, "sync 2: exits 0 (a rejected delivery does not fail the sync)", p.stderr[-800:])
    aud = dict(wh(conn, "select package, status || ':' || coalesce(rule_id,'') from tester.pf_package_audit"))
    R.check(aud.get("20240303/LB_ONC101_003.zip") == "rejected:package_checksum", "sync 2: bad checksum rejected with its rule id", aud)
    R.check(aud.get("20240304/LB_ONC101_004.zip") == "loaded:", "sync 2: next good delivery still loads", aud)
    n = wh(conn, "select count(*) from tester.lab_results")[0][0]
    R.check(n == 60, "sync 2: 6 new rows (visit 4), nothing from the rejected delivery", n)

    before = wh(conn, "select count(*) from tester.pf_package_audit")[0][0]
    p = tester(conn, cfg, state, "sftp-sync3", env)
    after = wh(conn, "select count(*) from tester.pf_package_audit")[0][0]
    R.check(p.returncode == 0 and before == after, "sync 3: no changes -> nothing re-emitted (idempotent)", (before, after))

    (root / "20240303" / "LB_ONC101_003.zip.sha256").write_text(f"{__import__('hashlib').sha256(d3).hexdigest()}  LB_ONC101_003.zip\n")
    os.utime(root / "20240303" / "LB_ONC101_003.zip.sha256", (OLD + 60, OLD + 60))  # re-uploaded later
    p = tester(conn, cfg, state, "sftp-sync4", env)
    st = wh(conn, "select status from tester.pf_package_audit where package = '20240303/LB_ONC101_003.zip'")
    R.check(("loaded",) in st, "sync 4: vendor fixed only the sidecar -> the rejected delivery is retried and loads", st)

    fx.put(root, "20240301/LB_ONC101_001.zip", fx.delivery(fx.rows(3, 1)))
    os.utime(root / "20240301" / "LB_ONC101_001.zip", (OLD, OLD))
    os.utime(root / "20240301" / "LB_ONC101_001.zip.sha256", (OLD, OLD))
    p = tester(conn, cfg, state, "sftp-sync5", env)
    st = wh(conn, "select status, rule_id from tester.pf_package_audit where package = '20240301/LB_ONC101_001.zip' order by finished_at")
    R.check(("rejected", "on_redelivery") in st, "sync 5: loaded delivery re-sent with different content is rejected (on_redelivery)", st)

    q = pf(repo, "dq", "run", "oncology_labs", "--engine", "duckdb", "--db", str(conn / "files" / "warehouse.db"), "--schema", "tester", env=env, ok=None)
    res = {r["id"]: r["status"] for r in q["data"].get("results", [])}
    R.check(res and all(v == "pass" for v in res.values()), "post-load quality checks pass on the tester warehouse", res or q["issues"])
    st_rows = wh(conn, "select count(*) from tester.pf_rule_results")[0][0]
    R.check(st_rows > 0, "rule results recorded for every delivery (audit evidence)", st_rows)
    leaked = wh(conn, "select count(*) from tester.pf_quarantine where cast(columns as varchar) like '%S0001%'")[0][0] if \
        wh(conn, "select count(*) from information_schema.columns where table_name='pf_quarantine' and column_name='columns'")[0][0] else 0
    R.check(leaked == 0, "quarantine evidence holds no data values", leaked)


# ------------------------------------------------------------------------------------------------ scenario: FTPS
FTP_HEADER = "VISIT_ID,SITE_NO,STUDY,VISIT_TYPE,VISIT_DATE,REPORT_DATE,OPEN_ISSUES,CRA_EMAIL"


def ftp_rows(start: int, n: int) -> list[str]:
    out = []
    for i in range(start, start + n):
        vt = ["SIV", "IMV", "COV"][i % 3]
        out.append(f"MV{i:05d},{1000 + i % 7},D1234C00001,{vt},2024-03-{1 + i % 20:02d},2024-03-{3 + i % 20:02d},{0 if vt == 'COV' else i % 4},cra{i % 5}@cro.example")
    return out


def ftp_file(lines: list[str]) -> bytes:
    return ("\n".join([FTP_HEADER, *lines]) + "\n").encode()


def scenario_ftp(state: Path, E: dict, env: dict) -> None:
    R.scenario = "ftp"
    print("\n== FTPS: site_monitoring (scaffold + requirement document intake + owner edits)", flush=True)
    repo = new_repo(state, env)
    out = pf(repo, "new", "site_monitoring", "--source", "ftp", "--destination", "s3",
             "--requirements", str(FRAMEWORK / "dryrun" / "requirements" / "site_monitoring.md"), env=env)
    intake = out["data"]["intake"]
    R.check(len(intake["applied"]) >= 10, "intake drafted values from the requirement document", len(intake["applied"]))
    R.check(len(intake["open_questions"]) > 0, "intake listed open questions for owners", intake["open_questions"][:3])
    R.check(all("evidence" in a or "entity" in a or "rule" in a for a in intake["applied"]), "every drafted value carries evidence")
    sttm = (repo / "pipelines" / "site_monitoring" / "sttm.yaml").read_text()
    R.check("kind: regex" in sttm and "kind: allowed_values" in sttm and "kind: range" in sttm,
            "single-value constraints became deterministic checks (regex, allowed_values, range)")
    R.check(sttm.count("kind: english") == 2, "only cross-column rules stayed in English (2)", sttm.count("kind: english"))
    v = pf(repo, "validate", "site_monitoring", env=env, ok=None)
    req = sorted({(i["file"], i["loc"]) for i in v["issues"] if i.get("code") == "required"})
    print(f"  owners still to fill: {len(req)}", flush=True)

    # owners fill what the document could not (governance, infrastructure) and answer the open questions
    owners = {("pipeline.yaml", "gxp_tier"): "non_gxp", ("pipeline.yaml", "owners.product_owner"): "@clinops-po",
              ("pipeline.yaml", "owners.data_engineer"): "@rd-data-eng", ("pipeline.yaml", "owners.governance"): "@rd-governance",
              ("pipeline.yaml", "owners.quality"): "@rd-data-quality", ("destination.yaml", "destination.processing_region"): "AWS_US_EAST_1",
              ("quality.yaml", "ticket_route"): "jira:CLINOPS", ("policy.yaml", "default_classification"): "internal",
              ("policy.yaml", "retention.days"): 2555, ("policy.yaml", "retention.rationale"): "7 years per business requirement",
              ("policy.yaml", "retention.object_lock_days"): 2555}
    for envn, acct in (("dev", "111111111111"), ("prod", "222222222222")):
        owners.update({("destination.yaml", f"destination.environments.{envn}.fivetran_group"): f"rd_lake_{envn}",
                       ("destination.yaml", f"destination.environments.{envn}.bucket"): f"az-rd-lake-{envn}",
                       ("destination.yaml", f"destination.environments.{envn}.region"): "us-east-1",
                       ("destination.yaml", f"destination.environments.{envn}.fivetran_role_arn"): f"arn:aws:iam::{acct}:role/fivetran-mdls",
                       ("policy.yaml", f"environments.{envn}.audit.bucket"): f"az-rd-pf-audit-{envn}",
                       ("policy.yaml", f"environments.{envn}.originals.bucket"): f"az-rd-originals-{envn}",
                       ("policy.yaml", f"environments.{envn}.originals.region"): "us-east-1"})
    fill(repo, "site_monitoring", owners, env)
    fill(repo, "site_monitoring", {("source.yaml", "source.host"): "127.0.0.1", ("source.yaml", "source.port"): E["ports"]["ftp"],
                                   ("source.yaml", "source.username"): E["ftp_user"],
                                   ("source.yaml", "source.ca_cert"): {"secret": "ftp_ca_bundle"},
                                   ("policy.yaml", "classification"): {"site_monitoring.cra_email": "confidential"}}, env)
    # the quality owner answers the questions intake raised by making the rule text explicit (the normal workflow)
    import yaml

    sm = yaml.safe_load((repo / "pipelines/site_monitoring/sttm.yaml").read_text())["entities"][0]
    eng = [c for c in sm.get("row_checks", []) if c["kind"] == "english"]
    rid_report = next(c["id"] for c in eng if "10" in c["text"])
    rid_cov = next(c["id"] for c in eng if "COV" in c["text"])
    attr_check = {a["name"]: (a.get("checks") or [{}])[0].get("id") for a in sm["attributes"]}
    answers = {
        rid_report:
            "When REPORT_DATE is present, it must be on or after VISIT_DATE and at most 10 calendar days after it "
            "(exactly 10 days passes). When REPORT_DATE is missing the report is not finalised yet and the row passes. "
            "When REPORT_DATE is present but VISIT_DATE is missing, the row fails.",
        rid_cov:
            "When VISIT_TYPE is exactly COV (case-sensitive), OPEN_ISSUES must be present and equal to 0. "
            "Rows with any other VISIT_TYPE, or with VISIT_TYPE missing, pass this rule.",
    }
    for rid, text in answers.items():
        r = pf(repo, "set", "site_monitoring", "sttm.yaml", f"entities[site_monitoring].row_checks[{rid}].text", json.dumps(text), env=env, ok=None)
        R.check(r["ok"], f"owner reworded {rid} to answer the open questions", r["issues"][:2])
    examples = {
        rid_report: [{"input": {"visit_date": "2024-03-01", "report_date": "2024-03-11"}, "expected": "pass"},
                     {"input": {"visit_date": "2024-03-01", "report_date": "2024-03-12"}, "expected": "fail"},
                     {"input": {"visit_date": "2024-03-01", "report_date": None}, "expected": "pass"}],
        rid_cov: [{"input": {"visit_type": "COV", "open_issues": 0}, "expected": "pass"},
                  {"input": {"visit_type": "COV", "open_issues": 2}, "expected": "fail"},
                  {"input": {"visit_type": "IMV", "open_issues": 5}, "expected": "pass"}],
    }
    for rid, ex in examples.items():
        r = pf(repo, "set", "site_monitoring", "sttm.yaml", f"entities[site_monitoring].row_checks[{rid}].examples", json.dumps(ex), env=env, ok=None)
        R.check(r["ok"], f"owner labelled 3 examples for {rid}", r["issues"][:2])
    v = pf(repo, "validate", "site_monitoring", env=env, ok=None)
    left = [(i["file"], i["loc"], i["msg"][:80]) for i in v["issues"] if i.get("severity") == "error"]
    R.check(not left, "spec complete after owner edits (pf set keeps comments, rolls back invalid edits)", left)
    bad_edit = pf(repo, "set", "site_monitoring", "source.yaml", "source.tls", json.dumps("yes"), env=env, ok=None)
    R.check(not bad_edit["ok"] and "tls: explicit" in (repo / "pipelines/site_monitoring/source.yaml").read_text(),
            "an invalid edit is refused and rolled back")

    c = pf(repo, "rules", "compile", "site_monitoring", env=env, ok=None)
    R.check(c["ok"], "English rules compiled to CEL with generated boundary cases", c["issues"][:3])
    agent_env = dict(env, PF_AGENT_SESSION="1")
    denied = pf(repo, "rules", "accept", "site_monitoring", rid_report, "--by", "@x", env=agent_env, ok=None)
    R.check(not denied["ok"] and denied["issues"][0]["code"] == "forbidden", "an agent session cannot accept a rule")
    data = c["data"]
    for rid, info in data.items():
        R.check(not info.get("ambiguities"), f"{rid}: no open questions left after the owner's rewording", info.get("ambiguities"))
        R.check(info.get("verification") == "passed", f"{rid}: every generated and owner case verified", info.get("verification"))
        a = pf(repo, "rules", "accept", "site_monitoring", rid, "--by", "@rd-data-quality", env=env, ok=None)
        R.check(a["ok"], f"owner accepted {rid}: {info.get('expression','')[:90]}", a["issues"][:2])

    fxd = repo / "pipelines" / "site_monitoring" / "fixtures"
    (fxd / "good").mkdir(parents=True, exist_ok=True)
    (fxd / "good" / "SMV_20240301_01.csv").write_bytes(ftp_file(ftp_rows(1, 30)))
    bad = fxd / "bad"

    def case(name: str, lines: list[str], expect: str, fname: str = "SMV_20240301_01.csv") -> None:
        (bad / name).mkdir(parents=True, exist_ok=True)
        (bad / name / fname).write_bytes(ftp_file(lines))
        (bad / name / "expect.yaml").write_text(expect)

    base = ftp_rows(1, 30)
    case("bad_site", [base[0].replace(",1001,", ",10X1,")] + base[1:], f"status: loaded\nquarantined: {{{attr_check['site_no']}: 1}}\n")
    case("bad_visit_type", [base[1].replace(",COV,", ",XYZ,").replace(",IMV,", ",XYZ,").replace(",SIV,", ",XYZ,")] + base[2:],
         f"status: loaded\nquarantined: {{{attr_check['visit_type']}: 1}}\n")
    neg = base[0].split(","); neg[6] = "-1"
    case("negative_issues", [",".join(neg)] + base[1:], f"status: loaded\nquarantined: {{{attr_check['open_issues']}: 1}}\n")
    late = base[0].split(","); late[4] = "2024-03-01"; late[5] = "2024-03-20"
    case("late_report", [",".join(late)] + base[1:], f"status: loaded\nquarantined: {{{rid_report}: 1}}\n")
    cov = base[1].split(","); cov[6] = "3"  # MV00002 is a COV visit
    case("cov_with_issues", [base[0], ",".join(cov)] + base[2:], f"status: loaded\nquarantined: {{{rid_cov}: 1}}\n")
    case("empty", [], "status: rejected\nrule_id: site_monitoring_has_rows\n")
    case("misnamed_file", base, "status: unmatched\nrule_id: package_match\n", fname="SMV_2024031_1.csv")
    b = pf(repo, "build", "site_monitoring", env=env, ok=None)
    R.check(b["ok"], "pf build", b["issues"][:3])
    t = pf(repo, "test", "site_monitoring", env=env, ok=None)
    R.check(t["ok"], "pf test on owner-written fixtures", [i["msg"][:300] for i in t["issues"]][:4])

    root = Path(E["ftp_root"]) / "exports" / "monitoring"
    root.mkdir(parents=True, exist_ok=True)
    for name, lines in (("SMV_20240301_01.csv", ftp_rows(1, 30)), ("SMV_20240301_02.csv", ftp_rows(25, 10)), ("SMV_20240302_01.csv", ftp_rows(40, 5))):
        (root / name).write_bytes(ftp_file(lines))
        os.utime(root / name, (OLD, OLD))
    (root / "SMV_20240303_01.csv").write_bytes(ftp_file(ftp_rows(60, 5)))  # too new: readiness must hold it back
    conn = repo / "pipelines" / "site_monitoring" / "build" / "connector"
    ca = Path(E["ftp_root"]).parent / "ftps.pem"
    pem = ca.read_text()
    cert = pem[pem.index("-----BEGIN CERTIFICATE-----"):]
    cfg = {"pf_environment": "dev", "ftp_password": E["ftp_password"], "ftp_ca_bundle": cert,
           "originals_access_key_id_dev": "testing", "originals_secret_access_key_dev": "testing", "pf_originals_endpoint_url": E["s3_endpoint"]}
    try:
        s3(E).create_bucket(Bucket="az-rd-originals-dev")
    except Exception:  # noqa: BLE001 - created by the sftp scenario
        pass
    p = tester(conn, cfg, state, "ftp-sync1", env)
    R.check(p.returncode == 0, "FTPS sync: exits 0 over verified TLS (pinned CA)", p.stdout[-600:] + p.stderr[-600:])
    n = wh(conn, "select count(*) from tester.site_monitoring")[0]
    R.check(n[0] == 39, "FTPS sync: 39 distinct visits (re-issue _02 overlapped 6 visits and updated them by primary key)", n)
    aud = wh(conn, "select package, status from tester.pf_package_audit order by package")
    R.check([a[0] for a in aud] == ["SMV_20240301_01.csv", "SMV_20240301_02.csv", "SMV_20240302_01.csv"],
            "FTPS sync: loaded in date, sequence order; the 2-minute-old file waits for readiness", aud)

    other = state / "other-ca.pem"
    subprocess.run(["openssl", "req", "-x509", "-newkey", "rsa:2048", "-nodes", "-days", "1", "-subj", "/CN=127.0.0.1",
                    "-addext", "subjectAltName=IP:127.0.0.1", "-keyout", str(state / "other-ca.key"), "-out", str(other)],
                   check=True, capture_output=True)
    p = tester(conn, dict(cfg, ftp_ca_bundle=other.read_text()), state, "ftp-untrusted", env)
    log = p.stdout + p.stderr
    R.check("not trusted" in log or "certificate" in log.lower(), "FTPS without the CA: the self-signed server is refused (no silent MITM)", log[-500:])


# ------------------------------------------------------------------------------------------------ scenario: RDS
def scenario_rds(state: Path, E: dict, env: dict) -> None:
    R.scenario = "rds"
    print("\n== RDS (PostgreSQL 16): trial_sites (incremental cursor, docx requirement)", flush=True)
    import psycopg2

    pg = E["pg"]
    admin = psycopg2.connect(host="127.0.0.1", port=pg["port"], user=pg["admin_user"], password=pg["admin_password"], dbname="postgres")
    admin.autocommit = True
    cur = admin.cursor()
    cur.execute("DROP DATABASE IF EXISTS ctms")
    cur.execute("CREATE DATABASE ctms")
    admin.close()
    db = psycopg2.connect(host="127.0.0.1", port=pg["port"], user=pg["admin_user"], password=pg["admin_password"], dbname="ctms")
    db.autocommit = True
    cur = db.cursor()
    cur.execute("""
        DROP ROLE IF EXISTS pf_reader;
        CREATE ROLE pf_reader LOGIN PASSWORD 'pf-reader-pw';
        CREATE SCHEMA clinical;
        CREATE TABLE clinical.sites (site_id varchar(4) PRIMARY KEY, site_name text NOT NULL, country char(2) NOT NULL,
            status text NOT NULL, activated_on date, enrolment_target int, updated_at timestamp NOT NULL);
        GRANT USAGE ON SCHEMA clinical TO pf_reader; GRANT SELECT ON clinical.sites TO pf_reader;""")
    for i in range(1, 251):
        cur.execute("INSERT INTO clinical.sites VALUES (%s,%s,%s,%s,%s,%s, timestamp '2024-03-01 00:00:00' + %s * interval '1 minute')",
                    (f"{i:04d}", f"Site {i}", ["GB", "SE", "US", "DE"][i % 4], ["ACTIVE", "CLOSED", "PLANNED"][i % 3],
                     None if i % 3 == 2 else f"2023-{1 + i % 12:02d}-15", 10 + i % 40, i))

    import docx

    d = docx.Document()
    d.add_heading("Requirement: Investigator sites from CTMS", 1)
    d.add_paragraph("Purpose: provide the current list of investigator sites for feasibility analytics.")
    d.add_paragraph("Source: the CTMS PostgreSQL database `ctms` (RDS), table clinical.sites. Read-only user pf_reader. "
                    "Rows change in place; updated_at is set on every change.")
    t = d.add_table(rows=1, cols=3)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Column", "Meaning", "Rules"
    for c in (("site_id", "Site number", "4 digits, unique"), ("site_name", "Site name", "mandatory"),
              ("country", "ISO country code", "two upper-case letters"), ("status", "ACTIVE, CLOSED or PLANNED", "mandatory"),
              ("activated_on", "Date the site was activated", "required when status is ACTIVE"),
              ("enrolment_target", "Planned subjects", "between 1 and 500"), ("updated_at", "Last change", "")):
        r = t.add_row().cells
        r[0].text, r[1].text, r[2].text = c
    doc = state / "trial_sites.docx"
    d.save(doc)

    repo = new_repo(state, env)
    out = pf(repo, "new", "trial_sites", "--source", "rds", "--destination", "s3", "--entities", "sites", "--requirements", str(doc), env=env)
    R.check(len(out["data"]["intake"]["applied"]) >= 5, "intake from a .docx requirement (tables included)", len(out["data"]["intake"]["applied"]))
    sttm = (repo / "pipelines" / "trial_sites" / "sttm.yaml").read_text()
    R.check("cursor_column: updated_at" in sttm, "intake found the incremental cursor column")
    owners = {("pipeline.yaml", "gxp_tier"): "non_gxp", ("pipeline.yaml", "owners.product_owner"): "@feasibility-po",
              ("pipeline.yaml", "owners.data_engineer"): "@rd-data-eng", ("pipeline.yaml", "owners.governance"): "@rd-governance",
              ("pipeline.yaml", "owners.quality"): "@rd-data-quality", ("destination.yaml", "destination.processing_region"): "AWS_US_EAST_1",
              ("quality.yaml", "ticket_route"): "jira:FEAS", ("policy.yaml", "default_classification"): "internal",
              ("policy.yaml", "retention.days"): 3650, ("policy.yaml", "retention.rationale"): "Reference data, 10 years",
              ("policy.yaml", "retention.object_lock_days"): 365,
              ("source.yaml", "source.host"): "127.0.0.1", ("source.yaml", "source.port"): pg["port"],
              ("source.yaml", "source.engine"): "postgres", ("source.yaml", "source.database"): "ctms",
              ("source.yaml", "source.username"): "pf_reader", ("source.yaml", "source.ssl_mode"): "disable",
              ("source.yaml", "source.batch_size"): 100}
    for envn, acct in (("dev", "111111111111"), ("prod", "222222222222")):
        owners.update({("destination.yaml", f"destination.environments.{envn}.fivetran_group"): f"rd_lake_{envn}",
                       ("destination.yaml", f"destination.environments.{envn}.bucket"): f"az-rd-lake-{envn}",
                       ("destination.yaml", f"destination.environments.{envn}.region"): "us-east-1",
                       ("destination.yaml", f"destination.environments.{envn}.fivetran_role_arn"): f"arn:aws:iam::{acct}:role/fivetran-mdls",
                       ("policy.yaml", f"environments.{envn}.audit.bucket"): f"az-rd-pf-audit-{envn}"})
    owners[("pipeline.yaml", "requirement.consumers")] = ["Feasibility analytics"]
    v0 = pf(repo, "validate", "trial_sites", env=env, ok=None)
    unfilled = {(i["file"], i["loc"].rsplit(".", 1)[0] if i["loc"].rsplit(".", 1)[-1].isdigit() else i["loc"])
                for i in v0["issues"] if i.get("code") == "required"}
    fill(repo, "trial_sites", {k: val for k, val in owners.items() if k in unfilled or k[0] == "source.yaml"}, env)
    v = pf(repo, "validate", "trial_sites", env=env, ok=None)
    R.check(any(i.get("loc") == "source.ssl_mode" for i in v["issues"]), "ssl_mode disable is flagged (dry run only)")
    left = [(i["file"], i["loc"], i["msg"][:100]) for i in v["issues"] if i.get("severity") == "error"]
    R.check(not left, "RDS spec complete after owner edits", left)
    if left:
        return
    import yaml

    ent = yaml.safe_load((repo / "pipelines/trial_sites/sttm.yaml").read_text())["entities"][0]
    eng = [(f"row_checks[{c['id']}]", c) for c in ent.get("row_checks", []) if c["kind"] == "english"] + \
          [(f"attributes[{a['name']}].checks[{c['id']}]", c) for a in ent["attributes"] for c in a.get("checks", []) if c["kind"] == "english"]
    for path, rule in eng:
        if "ACTIVE" in rule["text"]:
            text = ("When STATUS is exactly ACTIVE (upper case), ACTIVATED_ON must be present. "
                    "For any other STATUS, or a missing STATUS, ACTIVATED_ON is not checked.")
            ex = [{"input": {"status": "ACTIVE", "activated_on": "2023-01-15"}, "expected": "pass"},
                  {"input": {"status": "ACTIVE", "activated_on": None}, "expected": "fail"},
                  {"input": {"status": "PLANNED", "activated_on": None}, "expected": "pass"}]
            pf(repo, "set", "trial_sites", "sttm.yaml", f"entities[sites].{path}.text", json.dumps(text), env=env, ok=None)
            r = pf(repo, "set", "trial_sites", "sttm.yaml", f"entities[sites].{path}.examples", json.dumps(ex), env=env, ok=None)
            R.check(r["ok"], f"owner answered and labelled {rule['id']}", r["issues"][:2])
        else:
            R.check(False, f"unexpected English rule from intake (should be a deterministic check?): {rule['id']}: {rule['text']}")
    c = pf(repo, "rules", "compile", "trial_sites", env=env, ok=None)
    for rid, info in (c.get("data") or {}).items():
        R.check(not info.get("ambiguities"), f"{rid}: no open questions left", info.get("ambiguities"))
        a = pf(repo, "rules", "accept", "trial_sites", rid, "--by", "@rd-data-quality", env=env, ok=None)
        R.check(a["ok"], f"owner accepted {rid}: {info.get('expression','')[:90]}", a["issues"][:2])
    fxd = repo / "pipelines" / "trial_sites" / "fixtures" / "good"
    fxd.mkdir(parents=True, exist_ok=True)
    import io as _io
    buf = _io.StringIO()
    cur.copy_expert("COPY (SELECT site_id, site_name, country, status, activated_on, enrolment_target, updated_at FROM clinical.sites ORDER BY site_id LIMIT 20) TO STDOUT WITH CSV HEADER", buf)
    (fxd / "sites.csv").write_text(buf.getvalue())
    b = pf(repo, "build", "trial_sites", env=env, ok=None)
    R.check(b["ok"], "pf build", b["issues"][:3])
    t = pf(repo, "test", "trial_sites", env=env, ok=None)
    R.check(t["ok"], "pf test (RDS fixtures)", [i["msg"][:200] for i in t["issues"]][:3])
    conn = repo / "pipelines" / "trial_sites" / "build" / "connector"
    cfg = {"pf_environment": "dev", "rds_password": "pf-reader-pw", "rds_ca_bundle": "-----BEGIN CERTIFICATE-----\nunused\n-----END CERTIFICATE-----"}
    p = tester(conn, cfg, state, "rds-sync1", env)
    R.check(p.returncode == 0, "RDS sync 1: exits 0 (read-only session, server-side cursor)", p.stdout[-600:] + p.stderr[-600:])
    n = wh(conn, "select count(*) from tester.sites")[0][0]
    q = wh(conn, "select count(*) from tester.pf_quarantine")[0][0]
    R.check(n == 250 and q == 0, "RDS sync 1: all 250 rows loaded (seed data satisfies every rule)", (n, q))
    cur.execute("UPDATE clinical.sites SET status='CLOSED', updated_at = timestamp '2024-04-01' WHERE site_id IN ('0001','0002')")
    cur.execute("INSERT INTO clinical.sites VALUES ('0999','New site','GB','PLANNED',NULL,20, timestamp '2024-04-02')")
    p = tester(conn, cfg, state, "rds-sync2", env)
    st = dict(wh(conn, "select site_id, status from tester.sites where site_id in ('0001','0002','0999')"))
    R.check(p.returncode == 0 and st.get("0001") == "CLOSED" and "0999" in st, "RDS sync 2: changed and new rows picked up by the cursor", st)
    a = wh(conn, "select rows_read from tester.pf_entity_audit order by _fivetran_synced desc limit 1")
    R.check(a and a[0][0] <= 5, "RDS sync 2: incremental read (only rows at/after the cursor)", a)
    try:
        cur.execute("SET ROLE pf_reader")
        cur.execute("DELETE FROM clinical.sites WHERE site_id='0003'")
        R.check(False, "the reader role cannot write")
    except Exception:  # noqa: BLE001
        R.check(True, "the database role used by the connector cannot write")
    db.close()


# ------------------------------------------------------------------------------------------------ scenario: release + deploy
def scenario_deploy(state: Path, E: dict, env: dict) -> None:
    R.scenario = "deploy"
    print("\n== Release and Git-driven deploy (mock Fivetran API)", flush=True)
    import socket

    sys.path.insert(0, str(FRAMEWORK / "dryrun"))
    import mock_fivetran

    repo = new_repo(state, env)
    with socket.socket() as s_:
        s_.bind(("127.0.0.1", 0))
        port = s_.getsockname()[1]
    mock_fivetran.serve(state, port, ["rd_lake_dev", "rd_lake_prod"])
    sdk_cfg = state / "home" / ".fivetran" / "connector_sdk"
    sdk_cfg.mkdir(parents=True, exist_ok=True)
    (sdk_cfg / "_config.json").write_text(json.dumps({"production_base_url": f"http://127.0.0.1:{port}"}))
    key = base64.b64encode(mock_fivetran.KEY.encode()).decode()
    denv = dict(env, PF_FIVETRAN_API_BASE=f"http://127.0.0.1:{port}/v1", FIVETRAN_API_KEY=key,
                AWS_ENDPOINT_URL=E["s3_endpoint"], AWS_ACCESS_KEY_ID="testing", AWS_SECRET_ACCESS_KEY="testing", AWS_DEFAULT_REGION="us-east-1")
    name = "oncology_labs"
    run(["git", "add", "-A"], repo)
    run(["git", "commit", "-qm", "pipelines for dry run"], repo)
    chk = pf(repo, "release", name, "--check", env=denv, ok=None)
    R.check(not chk["ok"], "release --check fails until a release record is committed", chk["issues"][:1])
    rel = pf(repo, "release", name, env=denv, ok=None)
    R.check(rel["ok"], "pf release writes the immutable artefact and the release record", rel["issues"][:2])
    run(["git", "add", "-A"], repo)
    run(["git", "commit", "-qm", "release record"], repo)
    R.check(pf(repo, "release", name, "--check", env=denv, ok=None)["ok"], "release --check passes once the record is committed")
    s3(E).create_bucket(Bucket="az-rd-pf-artefacts")
    up = pf(repo, "release", name, "--upload", "s3://az-rd-pf-artefacts/releases", env=denv, ok=None)
    R.check(up["ok"], "artefact uploaded write-once", up["issues"][:2])
    up2 = pf(repo, "release", name, "--upload", "s3://az-rd-pf-artefacts/releases", env=denv, ok=None)
    R.check(up2["ok"], "re-upload of the same version is idempotent (same sha accepted)", up2["issues"][:2])

    sec = {"PF_SECRET_SFTP_PRIVATE_KEY": Path(E["client_key"]).read_text(), "PF_SECRET_SUBJECT_KEY_SALT": "dry-run-salt-1234567890",
           "PF_SECRET_ORIGINALS_ACCESS_KEY_ID_DEV": "AKIADRYRUN", "PF_SECRET_ORIGINALS_SECRET_ACCESS_KEY_DEV": "dry-run-secret-value-xyz"}
    agent = pf(repo, "deploy", name, "--env", "dev", env=dict(denv, **sec, PF_AGENT_SESSION="1"), ok=None)
    R.check(not agent["ok"] and agent["issues"][0]["code"] == "forbidden", "an agent session cannot deploy")
    miss = pf(repo, "deploy", name, "--env", "dev", env=denv, ok=None)
    R.check(not miss["ok"], "deploy refuses when a secret is missing", miss["issues"][:1])
    ed = pf(repo, "fivetran", "ensure-destination", name, "--env", "dev", env=denv, ok=None)
    R.check(ed["ok"], "destination created (managed_data_lake, Glue, Iceberg) through the API", ed["issues"][:2])
    dry = pf(repo, "deploy", name, "--env", "dev", "--dry-run", env=dict(denv, **sec), ok=None)
    shown = json.dumps(dry)
    R.check(dry["ok"] and "dry-run-salt" not in shown and "BEGIN OPENSSH" not in shown, "deploy --dry-run shows the command with secrets masked", dry["issues"][:2])
    dep = pf(repo, "deploy", name, "--env", "dev", env=dict(denv, **sec), ok=None)
    R.check(dep["ok"], "deploy: official `fivetran deploy` uploads the package and creates the connection; schedule applied", dep["issues"][:3])
    reqs = [json.loads(line) for line in (state / "fivetran-requests.jsonl").read_text().splitlines()]
    R.check(any(r["path"].startswith("/v1/connector-sdk/packages") and r["method"] == "POST" for r in reqs), "package upload seen by the API")
    conn_req = [r for r in reqs if r["method"] == "POST" and r["path"] in ("/v1/connectors", "/v1/connections")]
    sl = conn_req[-1].get("json", {}).get("config", {}).get("secrets_list", []) if conn_req else []
    keys = {i.get("key"): i.get("value") for i in sl} if isinstance(sl, list) else {}
    R.check(keys.get("pf_environment") == "dev", "connection configuration (secrets_list) carries pf_environment=dev", sorted(keys))
    R.check(keys and not any(k.endswith("_prod") for k in keys), "dev connection receives only dev secrets", sorted(keys))
    R.check(conn_req and conn_req[-1]["json"].get("paused") is True, "new connection is created paused (SDK default)")
    sched = [r for r in reqs if r["method"] == "PATCH" and "/connections/" in r["path"]]
    R.check(sched and sched[-1]["json"].get("daily_sync_time") == "02:00" and sched[-1]["json"].get("paused") is True,
            "schedule from pipeline.yaml applied (dev: daily 02:00, paused)", sched[-1:] if sched else reqs[-3:])
    dr = pf(repo, "fivetran", "drift", name, "--env", "dev", env=denv, ok=None)
    R.check(dr["ok"] and not dr["data"].get("drift"), "no drift right after deploy", dr["data"])
    wf = (repo / ".github" / "workflows" / "pf-oncology-labs.yml").read_text()
    R.check("environment: prod" in wf and "pf deploy oncology_labs --env prod" in wf, "generated GitHub workflow deploys per environment gate")
    co = (repo / ".github" / "CODEOWNERS").read_text()
    R.check("@rd-governance" in co and "policy.yaml" in co, "CODEOWNERS routes policy changes to governance")


# ------------------------------------------------------------------------------------------------ failure drills
def scenario_drills(state: Path, E: dict, env: dict) -> None:
    R.scenario = "drills"
    print("\n== Failure drills (runtime against the real servers)", flush=True)
    repo = state / "repo"
    conn = repo / "pipelines" / "oncology_labs" / "build" / "connector"
    sys.path.insert(0, str(FRAMEWORK / "src"))
    from pf.testing import NullLog, RecorderOp
    from pf_runtime import engine

    lock = json.loads((conn / "spec.lock.json").read_text())
    base = {"pf_environment": "dev", "sftp_private_key": Path(E["client_key"]).read_text(), "subject_key_salt": "dry-run-salt-1234567890",
            "originals_access_key_id_dev": "testing", "originals_secret_access_key_dev": "testing", "pf_originals_endpoint_url": E["s3_endpoint"]}

    def go(lk, cfg, state_in=None):
        op = RecorderOp()
        engine.update(lk, cfg, state_in or {}, op, NullLog(), run_id="drill", sleep=lambda s: None)
        return op

    lk = json.loads(json.dumps(lock))
    lk["source"]["host_key"] = "ssh-ed25519 AAAAC3NzaC1lZDI1NTE5AAAAIOMqqnkVzrm0SdG6UOoqKLsabgH5C9okWi0dh2l9GKJl"
    op = go(lk, base)
    R.check(op.errors and "host key" in op.errors[0].lower() and not op.rows.get("lab_results"),
            "wrong pinned host key: sync stops before any data (MITM protection)", op.errors[:1])
    op = go(lock, dict(base, sftp_private_key="-----BEGIN OPENSSH PRIVATE KEY-----\nAAAA\n-----END OPENSSH PRIVATE KEY-----\n"))
    R.check(bool(op.errors) and not op.rows.get("lab_results"), "bad private key: configuration error, no data", op.errors[:1])
    lk = json.loads(json.dumps(lock))
    lk["source"]["port"] = 1
    t0 = time.time()
    op = go(lk, base)
    R.check(bool(op.errors) and "connect" in op.errors[0].lower(), "SFTP down: retried, then the sync fails loudly", op.errors[:1])
    op = go(lock, dict(base, pf_originals_endpoint_url="http://127.0.0.1:1"))
    R.check(bool(op.errors) and not op.rows.get("lab_results"), "originals store down: nothing is loaded without its archived original", op.errors[:1])
    op = go(lock, {k: v for k, v in base.items() if k != "subject_key_salt"})
    R.check(bool(op.errors) and "subject_key_salt" in op.errors[0], "missing pseudonymisation key: sync refuses to start", op.errors[:1])
    op = go(lock, dict(base, pf_environment="qa"))
    R.check(bool(op.errors), "unknown environment in configuration is refused", op.errors[:1])
    big = json.loads(json.dumps(lock))
    big["source"]["package"]["max_package_bytes"] = 100
    op = go(big, base)
    aud = [(r["status"], r["rule_id"]) for r in op.rows.get("pf_package_audit", [])]
    R.check(aud and all(a == ("rejected", "max_package_bytes") for a in aud), "oversized delivery rejected before download", aud[:2])
    del t0


SCENARIOS = {"sftp": scenario_sftp, "ftp": scenario_ftp, "rds": scenario_rds, "deploy": scenario_deploy, "drills": scenario_drills}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("state")
    ap.add_argument("--llm", default="replay", choices=["replay", "record", "agent-sdk"])
    ap.add_argument("--only", default=",".join(SCENARIOS))
    a = ap.parse_args()
    state = Path(a.state).resolve()
    E = json.loads((state / "env.json").read_text())
    home = state / "home"
    (home / ".fivetran" / "connector_sdk").mkdir(parents=True, exist_ok=True)
    real_tester = Path.home() / ".fivetran" / "connector_sdk" / "tester"
    link = home / ".fivetran" / "connector_sdk" / "tester"
    if real_tester.exists() and not link.exists():
        link.symlink_to(real_tester)
    env = base_env(state, a.llm)
    for name in a.only.split(","):
        try:
            SCENARIOS[name](state, E, env)
        except Exception as exc:  # noqa: BLE001 - a crashed scenario is a failed check, the run continues
            R.scenario = name
            R.check(False, f"scenario {name} crashed", f"{type(exc).__name__}: {exc}\n{traceback.format_exc()[-1500:]}")
    failed = [c for c in R.checks if not c["ok"]]
    (state / "DRY-RUN-REPORT.json").write_text(json.dumps({"passed": len(R.checks) - len(failed), "failed": len(failed), "checks": R.checks}, indent=2, default=str))
    print(f"\n{len(R.checks) - len(failed)} passed, {len(failed)} failed", flush=True)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
